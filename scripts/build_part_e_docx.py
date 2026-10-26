"""Render report/part-e.md into the Part E .docx and its PDF.

The report is authored in Markdown (`report/part-e.md`, produced by
`scripts/build_part_e_md.py`). This script is the only step that produces Word:
it opens the unit's template, keeps its styles (Times New Roman 12 pt body,
24/18/14 pt bold headings, 2.54 cm margins), and emits the Markdown blocks with
that typography — plus the static contents page and the page-number footer,
which need the rendered PDF's page map.

    python scripts/build_part_e_docx.py
"""
from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))

import md_blocks  # noqa: E402
from report_layout import Layout, add_page_footer, force_font  # noqa: E402

REPO = Path(__file__).resolve().parents[1]
MD = REPO / "report" / "part-e.md"
OUT = REPO / "report" / "RailTech-A3-PartE-Report.docx"
TEMPLATE = REPO / "templates" / "part-e-template-2.docx"
PAGES = REPO / "report" / ".toc-pages.json"

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
TEXT_WIDTH_CM = 15.9
HEADING_PT = {1: 24, 2: 18, 3: 14, 4: 12}


def clear_body(doc: Document) -> None:
    """Keep the section properties, drop the template's own content."""
    body = doc.element.body
    for child in list(body):
        if child.tag.endswith("}sectPr"):
            continue
        body.remove(child)


def heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(HEADING_PT.get(level, 12))
    run.font.name = "Times New Roman"
    force_font(run._r, "Times New Roman")


def para(doc: Document, text: str, size: float = 12) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    for chunk, bold, italic in md_blocks.inline(text):
        run = p.add_run(chunk)
        run.bold = bold or None
        run.italic = italic or None
        run.font.size = Pt(size)
        run.font.name = "Times New Roman"
        force_font(run._r, "Times New Roman")


def cover_para(doc: Document, text: str, index: int) -> None:
    """The title page: centred, with sizes taken from the template's hierarchy."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10 if index < 2 else 6)
    size, bold = 12, False
    if index == 0:
        size, bold = 16, True
        p.paragraph_format.space_before = Pt(90)
    elif index == 1:
        size, bold = 14, True
    elif text.startswith("**Project Title"):
        bold = True
    for chunk, b, i_ in md_blocks.inline(text):
        run = p.add_run(chunk)
        run.bold = b or bold or None
        run.italic = i_ or None
        run.font.size = Pt(size)
        run.font.name = "Times New Roman"
        force_font(run._r, "Times New Roman")


def caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(10)
    run.font.name = "Times New Roman"
    force_font(run._r, "Times New Roman")


def table(doc: Document, block: md_blocks.Block) -> None:
    rows = block.rows
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "808080")
        borders.append(el)
    t._tbl.tblPr.append(borders)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r in t.rows:
        r._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            if block.align and j < len(block.align) and block.align[j] == "r":
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = p.add_run(value)
            run.font.size = Pt(9)
            run.font.name = "Times New Roman"
            force_font(run._r, "Times New Roman")
            if i == 0:
                run.bold = True
                tcPr = cell._tc.get_or_add_tcPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), "EFEFEF")
                tcPr.append(shd)


def figure(doc: Document, block: md_blocks.Block) -> None:
    path = (MD.parent / block.src).resolve()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Cm(TEXT_WIDTH_CM * block.width))


def toc(doc: Document, pages: dict[str, int] | None, headings: list[str]) -> None:
    for name in headings:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.tab_stops.add_tab_stop(
            Cm(TEXT_WIDTH_CM - 0.3), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
        )
        run = p.add_run(name)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        p.add_run("\t")
        num = p.add_run(str(pages.get(name, "")) if pages else "")
        num.font.name = "Times New Roman"
        num.font.size = Pt(12)


def render(pages: dict[str, int] | None) -> None:
    doc = Document(str(TEMPLATE))
    clear_body(doc)

    blocks = md_blocks.load(MD)
    headings = [b.text for b in blocks
                if b.kind == "heading" and b.text != "Table of Contents"]
    in_cover = True
    cover_idx = 0
    i = 0
    while i < len(blocks):
        b = blocks[i]
        if b.kind == "break":
            p = doc.add_paragraph()
            p.add_run().add_break(WD_BREAK.PAGE)
            in_cover = False
        elif b.kind == "para" and b.text.strip() == "@@TOC@@":
            toc(doc, pages, headings)
        elif b.kind == "para" and in_cover:
            cover_para(doc, b.text, cover_idx)
            cover_idx += 1
        elif b.kind == "heading":
            heading(doc, b.text, b.level)
        elif b.kind == "para":
            para(doc, b.text)
        elif b.kind == "bullet":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.65)
            p.paragraph_format.first_line_indent = Cm(-0.4)
            p.paragraph_format.space_after = Pt(3)
            for chunk, bold, italic in md_blocks.inline("•  " + b.text):
                run = p.add_run(chunk)
                run.bold = bold or None
                run.italic = italic or None
                run.font.size = Pt(12)
                run.font.name = "Times New Roman"
                force_font(run._r, "Times New Roman")
        elif b.kind == "table":
            table(doc, b)
        elif b.kind == "figure":
            figure(doc, b)
            if i + 1 < len(blocks) and blocks[i + 1].kind == "caption":
                caption(doc, blocks[i + 1].text)
                i += 1
            else:
                caption(doc, b.text)
        elif b.kind == "caption":
            caption(doc, b.text)
        i += 1

    # the @@TOC@@ marker arrives as a para block: render it in place
    add_page_footer(doc, "RailTech, 43008 Assessment 3 (Part E)",
                    font="Times New Roman", size=9)
    doc.save(str(OUT))


def pdf_pages() -> dict[str, int]:
    pdf = OUT.with_suffix(".pdf")
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(pdf.parent), str(OUT)],
        capture_output=True, check=True,
    )
    text = subprocess.run(
        ["pdftotext", "-layout", str(pdf), "-"], capture_output=True, text=True, check=True
    ).stdout
    headings = [b.text for b in md_blocks.load(MD)
                if b.kind == "heading" and b.text != "Table of Contents"]
    found: dict[str, int] = {}
    for i, page in enumerate(text.split("\f"), 1):
        lines = [ln.strip() for ln in page.splitlines() if "....." not in ln]
        for name in headings:
            if name not in found and any(ln.startswith(name) for ln in lines):
                found[name] = i
    return found


def main() -> int:
    pages = json.loads(PAGES.read_text()) if PAGES.exists() else None
    render(pages)
    found = pdf_pages()
    if pages != found:
        PAGES.write_text(json.dumps(found, indent=1))
        render(found)
        found = pdf_pages()
    print("wrote", OUT, "| pages:", max(found.values()) if found else "?")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
