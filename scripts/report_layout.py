"""Shared layout helpers for the two report builders.

`build_report.py` writes the standalone 15-page report; `build_part_e_report.py`
writes the same content into the unit's Part E template.  Both need the same
paragraph, table, figure and caption styling, so it lives here rather than in
each builder.
"""
from __future__ import annotations

from pathlib import Path

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = RGBColor(0x1F, 0x38, 0x64)
GREY = RGBColor(0x40, 0x40, 0x40)
FILL_HEAD = "1F3864"
FILL_TOTAL = "DCE6F2"
FILL_BAND = "F2F5FA"


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
THEME_ATTRS = ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme")


def force_font(element, name: str) -> None:
    """Set the four rFonts slots and drop any theme reference that would win."""
    rPr = element.find(f"{W}rPr")
    if rPr is None:
        return
    rFonts = rPr.find(f"{W}rFonts")
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in THEME_ATTRS:
        key = f"{W}{attr}"
        if key in rFonts.attrib:
            del rFonts.attrib[key]
    for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
        rFonts.set(f"{W}{slot}", name)


def style_document(doc, heading_colour=NAVY, body_size=11,
                   margins=(2.2, 2.3), line_spacing=1.15) -> None:
    """Calibri body text, navy headings, A4 page with the given margins."""
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin = sec.bottom_margin = Cm(margins[0])
    sec.left_margin = sec.right_margin = Cm(margins[1])

    normal = doc.styles["Normal"]
    force_font(normal.element, "Calibri")
    normal.font.name = "Calibri"
    normal.font.size = Pt(body_size)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = line_spacing
    for name, size, before in (("Heading 1", 15, 18), ("Heading 2", 12.5, 14)):
        st = doc.styles[name]
        force_font(st.element, "Calibri")
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = heading_colour
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(8)
        st.paragraph_format.keep_with_next = True


def shade(cell, fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


class Layout:
    """Paragraph/table/figure emitters bound to one Document."""

    def __init__(self, doc, fig_dir: Path, max_width_cm: float | None = None,
                 font: str | None = None):
        self.doc = doc
        self.fig_dir = Path(fig_dir)
        self.max_width_cm = max_width_cm
        self.font = font

    def _style_run(self, run):
        """Pin the run font when this document has its own typeface."""
        if self.font:
            run.font.name = self.font
            force_font(run._r, self.font)
        return run

    def body(self, text: str):
        p = self.doc.add_paragraph()
        self._style_run(p.add_run(text))
        p.paragraph_format.space_after = Pt(6)
        return p

    def bullet(self, text: str):
        if "List Bullet" in {s.name for s in self.doc.styles}:
            p = self.doc.add_paragraph(style="List Bullet")
            self._style_run(p.add_run(text))
        else:
            # the Part E template ships no list style: draw the bullet ourselves
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.65)
            p.paragraph_format.first_line_indent = Cm(-0.4)
            self._style_run(p.add_run("\u2022  " + text))
        p.paragraph_format.space_after = Pt(3)
        return p

    def caption(self, text: str):
        p = self.doc.add_paragraph()
        r = self._style_run(p.add_run(text))
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = GREY
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(12)
        return p

    def figure(self, name: str, width_cm: float, text: str, max_width_cm: float | None = None):
        max_width_cm = max_width_cm or self.max_width_cm
        if max_width_cm:
            width_cm = min(width_cm, max_width_cm)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(str(self.fig_dir / name), width=Cm(width_cm))
        self.caption(text)
        return p

    def table(self, rows, widths, text: str, max_width_cm: float | None = None):
        max_width_cm = max_width_cm or self.max_width_cm
        if max_width_cm and sum(widths) > max_width_cm:
            scale = max_width_cm / sum(widths)
            widths = [w * scale for w in widths]
        t = self.doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for row in t.rows:
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for i, row in enumerate(rows):
            for j, value in enumerate(row):
                cell = t.cell(i, j)
                cell.width = Cm(widths[j])
                cell.text = ""
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                r = self._style_run(p.add_run(str(value)))
                r.font.size = Pt(9)
                if i == 0:
                    r.bold = True
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    shade(cell, FILL_HEAD)
                elif str(rows[i][0]) == "Total":
                    r.bold = True
                    shade(cell, FILL_TOTAL)
                elif i % 2 == 0:
                    shade(cell, FILL_BAND)
        self.caption(text)
        return t

    def page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_page_footer(doc, left_text: str, font: str | None = None, size: float = 8.5) -> None:
    """Centre footer: '<left_text>    |    page <n>' with a live PAGE field.

    The template ships no Footer style, so the paragraph is pinned to Normal;
    otherwise the footer falls back to the 11 pt Calibri default.
    """
    footer = doc.sections[0].footer.paragraphs[0]
    footer.style = doc.styles["Normal"]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def styled(text: str | None = None):
        run = footer.add_run(text or "")
        run.font.size = Pt(size)
        run.font.color.rgb = GREY
        if font:
            run.font.name = font
            force_font(run._r, font)
        return run

    styled(f"{left_text}    |    page ")

    for kind, value in (("begin", None), (None, " PAGE "), ("separate", None)):
        run = styled()
        if kind:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), kind)
            run._r.append(fld)
        else:
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = value
            run._r.append(instr)

    styled("1")

    run = styled()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "end")
    run._r.append(fld)


def set_document_defaults(doc, font: str, size_pt: float) -> None:
    """Pin the document defaults (the layer the template's duplicate styles fall to)."""
    root = doc.styles.element

    # the template carries duplicate styleIds; LibreOffice keeps the last one,
    # python-docx edits the first, so drop the emptier duplicates first
    seen: dict[str, object] = {}
    for st in list(root.findall(f"{W}style")):
        sid = st.get(f"{W}styleId")
        if not sid:
            continue
        if sid not in seen:
            seen[sid] = st
        else:
            keep, drop = (st, seen[sid]) if len(st) > len(seen[sid]) else (seen[sid], st)
            root.remove(drop)
            seen[sid] = keep

    rPr = root.find(f"{W}docDefaults/{W}rPrDefault/{W}rPr")
    if rPr is None:
        return
    rFonts = rPr.find(f"{W}rFonts")
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in THEME_ATTRS:
        key = f"{W}{attr}"
        if key in rFonts.attrib:
            del rFonts.attrib[key]
    for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
        rFonts.set(f"{W}{slot}", font)
    sz = rPr.find(f"{W}sz")
    if sz is None:
        sz = OxmlElement("w:sz")
        rPr.append(sz)
    sz.set(f"{W}val", str(int(size_pt * 2)))
