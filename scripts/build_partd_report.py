#!/usr/bin/env python3
"""Build the Assignment-3 Part-D (Environment/Interface Design) report.

Writes the completed report as a .docx (matching the university's Part-D
template and the style of Parts A-C) and converts it to PDF with LibreOffice.

The figures are produced by ``scripts/render_partd_figures.py``; run that first.

    python scripts/render_partd_figures.py
    python scripts/build_partd_report.py

Output (default):
    ~/Downloads/Assignment-3-PartD-Daniel-James-Martirosov-24948933.docx
    ~/Downloads/Assignment-3-PartD-Daniel-James-Martirosov-24948933.pdf
"""

from __future__ import annotations

import argparse
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

REPO_ROOT = Path(__file__).resolve().parents[1]
DOCS = REPO_ROOT / "docs"
DEFAULT_OUT = Path.home() / "Downloads"
BASENAME = "Assignment-3-PartD-Daniel-James-Martirosov-24948933"

MEMBERS = [
    ("Daniel James Martirosov", "24948933", "daniel.j.martirosov@student.uts.edu.au"),
    ("Laurean Lim", "25049863", "laurean.lim@student.uts.edu.au"),
    ("Ziheng Deng", "25595315", "Ziheng.Deng-1@student.uts.edu.au"),
]

FIGURES = [
    (
        "partd_scenario1_central_strathfield.png",
        16.5,
        "Figure 1: Scenario 1 - Central-Strathfield (92 x 37), mid-simulation under normal "
        "conditions. Each coloured square is one of the 74 TfNSW services (T1, T2, T3, T9, "
        "CCN, BMT) moving through the multi-track Main Suburban corridor; the red-roof "
        "markers are the stations and the grey grid is the real turnback track layout. "
        "Flatland's built-in renderer draws the whole corridor in one view, which is what "
        "makes the parallel-track rerouting problem visible.",
    ),
    (
        "partd_scenario1_corridor_zoom.png",
        16.5,
        "Figure 2: Scenario 1 zoomed on a station area, showing trains (coloured) occupying "
        "adjacent parallel tracks and platforms. Because each train is drawn in its own "
        "colour and every agent has a target marker, the interface makes it easy to see "
        "where conflicts and late-running build up in the corridor.",
    ),
    (
        "partd_scenario1_delayed_malfunction.png",
        16.5,
        "Figure 3: Scenario 1 under the delayed operating condition. Trains held by an "
        "injected malfunction are marked with a black cross (x) - they are stopped in place "
        "for 20-50 steps, exactly as track faults, maintenance and weather holds are modelled. "
        "This is the disruption view the trained policies are evaluated against.",
    ),
    (
        "partd_scenario2_granville_harris_park.png",
        12.5,
        "Figure 4: Scenario 2 - Granville-Harris Park (20 x 20), the Main Western junction "
        "carrying 126 services (T1 on the Main pair, T2 on the Suburban pair, T5 on the Y-link). "
        "The Old Main South from Granville and the Y-link from Harris Park merge onto a two-track "
        "branch down to Merrylands. Only a few services are on the map at this instant (the real "
        "timetable is sparse across the 16:00-19:00 window).",
    ),
    (
        "partd_scenario2_junction_zoom.png",
        16.0,
        "Figure 5: Scenario 2 zoom of the junction, showing the Old Main South and the Y-link "
        "converging onto the branch to Merrylands, with trains on the branch and the "
        "main/suburban tracks. Unlike Scenario 1's parallel-track rerouting, here the agents "
        "must manage right-of-way, waiting and queue propagation at a convergence.",
    ),
]


def _setup(doc: Document) -> None:
    # match the assignment template / Parts A-C: Times New Roman, body 12pt,
    # A4 page with 1 inch (2.54 cm) margins.
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), "Times New Roman")

    section = doc.sections[0]
    section.page_width = Cm(21.0)      # A4
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.54)     # 1 inch
    section.right_margin = Cm(2.54)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)


def _title(doc: Document) -> None:
    for line in ("43008: Reinforcement Learning",
                 "Assignment 3, Part-D",
                 "Environment/Interface Design"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(line)
        run.bold = True
        run.font.size = Pt(18)


def _field(doc: Document, label: str, value: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(f"{label} ")
    r.bold = True
    r.font.size = Pt(14)
    v = p.add_run(value)
    v.font.size = Pt(12)


def _members_table(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for cell, text in zip(hdr, ("Student Name", "Student ID", "Email ID")):
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.size = Pt(12)
    for name, sid, email in MEMBERS:
        cells = table.add_row().cells
        for cell, text in zip(cells, (name, sid, email)):
            run = cell.paragraphs[0].add_run(text)
            run.font.size = Pt(12)


def _heading(doc: Document, text: str, size: int = 14) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)


def _body(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(12)


def _bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        run.font.size = Pt(12)


def _figure(doc: Document, filename: str, width_cm: float, caption: str) -> None:
    path = DOCS / filename
    if not path.exists():
        raise FileNotFoundError(f"missing figure: {path} (run scripts/render_partd_figures.py)")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True   # keep the image with its caption
    p.add_run().add_picture(str(path), width=Cm(width_cm))

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_together = True
    run = cap.add_run(caption)
    run.italic = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)


def build(out_dir: Path) -> Path:
    doc = Document()
    _setup(doc)
    _title(doc)

    _field(doc, "Project Title and number (Refer to Dashboard):",
           "Train Schedule Optimisation under Disruptions (Project #27)")
    _field(doc, "Team Name:", "RailTech")
    doc.add_paragraph()
    p = doc.add_paragraph()
    run = p.add_run("Team Members:")
    run.bold = True
    run.font.size = Pt(14)
    _members_table(doc)
    doc.add_paragraph()

    _heading(doc, "Environment chosen:")
    _body(
        doc,
        "Our project uses Flatland (the open-source flatland-rl package, v4.3.0) as the "
        "reinforcement-learning environment. Flatland is a gridworld toolkit for developing "
        "and comparing multi-agent RL algorithms on railway networks; it was introduced "
        "through the NeurIPS 2020 Flatland Challenge and developed with the Swiss Federal "
        "Railways (SBB). It exposes a Gym-style environment (RailEnv) in which every train is "
        "an independent agent that has to run to a timetable, avoid collisions and cope with "
        "disruptions. We selected it because it gives a realistic multi-agent scheduling and "
        "conflict problem with built-in delay metrics, stochastic malfunctions and a ready "
        "renderer, and because it integrates cleanly with PyTorch and Stable-Baselines3 "
        "(PPO, A2C, DQN).",
    )
    _body(
        doc,
        "Instead of Flatland's generic random generator we designed two custom scenarios "
        "that recreate real Sydney corridors from Transport for NSW (TfNSW) GTFS timetable "
        "data for Thursday 15 October 2026, 16:00-19:00:",
    )
    _bullets(
        doc,
        [
            "Scenario 1 - Central-Strathfield (92 x 37 grid, 74 services from lines T1, T2, T3, "
            "T9, CCN and BMT): the busiest section of the Main Suburban corridor. It is a "
            "multi-track, parallel-track environment used to study rerouting under disruptions.",
            "Scenario 2 - Granville-Harris Park (20 x 20 grid, 126 services from lines T1, T2 and "
            "T5): a triangular junction on the Main Western line where the Old Main South (from "
            "Granville) and the Cumberland Y-link (from Harris Park) merge and run down to "
            "Merrylands. This is a junction/merge/branch-conflict environment used to study "
            "right-of-way, waiting and queue propagation.",
        ],
    )
    _body(
        doc,
        "Each train is an agent with a start position, heading, target and per-stop timetable; "
        "at each 20-second step it chooses one of Flatland's five discrete actions (do nothing, "
        "move left, forward, move right, stop). All agents share a single policy network, so "
        "experience from many trains trains one model. Under the normal condition the "
        "environment is deterministic; under the delayed condition disruptions are injected "
        "through Flatland's effects generator as Poisson malfunctions (rate 1/540, 20-50 step "
        "holds), modelling track faults, maintenance and weather holds. Observations, reward "
        "shaping (progress, punctuality and conflict penalties) and evaluation follow Parts B "
        "and C of this assignment.",
    )

    _heading(doc, "GUI/Environment Type:")
    _body(
        doc,
        "Custom environment (Flatland gridworld). We use Flatland's built-in graphical renderer "
        "as the interface, which the assignment specification permits in place of a separate "
        "GUI. The renderer draws the rail grid, the coloured trains (one colour per agent), the "
        "station/city markers and each agent's target marker. Two backends are used: an "
        "interactive SVG/pygame window to watch and demonstrate the trained policy live, and a "
        "headless PIL/SVG backend to produce reproducible still frames (and GIF rollouts) for "
        "this report and the demo video. Because the renderer is part of the environment, every "
        "screenshot below is generated directly from the same RailEnv instance the agents are "
        "trained and evaluated on - no separate interface or dashboard is needed.",
    )

    _heading(doc, "Screen shot(s) of the actual GUI/Environment:")
    _body(
        doc,
        "Figures 1-3 show Scenario 1 (Central-Strathfield) and Figures 4-5 show Scenario 2 "
        "(Granville-Harris Park). All frames are headless renders of the real environments.",
    )
    for filename, width, caption in FIGURES:
        _figure(doc, filename, width, caption)

    out_dir.mkdir(parents=True, exist_ok=True)
    docx_path = out_dir / f"{BASENAME}.docx"
    doc.save(docx_path)
    print(f"wrote {docx_path}")
    return docx_path


def to_pdf(docx_path: Path) -> Path | None:
    try:
        subprocess.run(
            ["soffice", "--headless", "--convert-to", "pdf", "--outdir",
             str(docx_path.parent), str(docx_path)],
            check=True,
            capture_output=True,
            timeout=180,
        )
    except (subprocess.CalledProcessError, FileNotFoundError, subprocess.TimeoutExpired) as exc:
        print(f"PDF conversion skipped: {exc}")
        return None
    pdf_path = docx_path.with_suffix(".pdf")
    print(f"wrote {pdf_path}")
    return pdf_path


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--out-dir", type=Path, default=DEFAULT_OUT)
    ap.add_argument("--no-pdf", action="store_true")
    args = ap.parse_args()

    docx_path = build(args.out_dir)
    if not args.no_pdf:
        to_pdf(docx_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
